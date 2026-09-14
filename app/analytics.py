from io import BytesIO
from pathlib import Path

from flask import Blueprint, abort, render_template, request, send_file
from flask_login import current_user
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from sqlalchemy import case, func, select

from app.access import (
    ensure_department_access,
    get_scoped_department_id,
)
from app.auth import (
    ROLE_ADMIN,
    ROLE_DEPARTMENT_HEAD,
    roles_required,
)
from app.extensions import db
from app.i18n import current_language, translate
from app.models import (
    Department,
    Plan,
    PlanItem,
    Publication,
    ReferenceValue,
)


ANALYTICS_PLAN_STATUS = "APPROVED"
ANALYTICS_FACT_STATUS = "APPROVED"
PUBLICATION_TYPE_REFERENCE = "PUBLICATION_TYPE"


analytics_bp = Blueprint(
    "analytics",
    __name__,
    url_prefix="/analytics",
)


def calculate_percentage(planned_count, actual_count):
    if planned_count <= 0:
        return 0.0

    return round(
        min(
            actual_count / planned_count * 100,
            100.0,
        ),
        2,
    )


def parse_optional_int(name):
    value = request.args.get(
        name,
        "",
    ).strip()

    if not value:
        return None

    try:
        parsed_value = int(value)
    except ValueError:
        abort(400)

    if parsed_value < 0:
        abort(400)

    return parsed_value


def get_filters():
    year = parse_optional_int("year")
    requested_department_id = parse_optional_int(
        "department_id"
    )
    publication_type_id = parse_optional_int(
        "publication_type_id"
    )

    ensure_department_access(
        requested_department_id
    )

    department_id = get_scoped_department_id(
        requested_department_id
    )

    return {
        "year": year,
        "department_id": department_id,
        "publication_type_id": publication_type_id,
    }


def get_years():
    return db.session.scalars(
        select(Plan.year)
        .where(
            Plan.status == ANALYTICS_PLAN_STATUS,
        )
        .distinct()
        .order_by(
            Plan.year.desc(),
        )
    ).all()


def get_departments():
    query = (
        select(Department)
        .where(
            Department.is_active.is_(True),
        )
        .order_by(
            Department.name,
        )
    )

    if current_user.role == ROLE_DEPARTMENT_HEAD:
        if current_user.department_id is None:
            abort(403)

        query = query.where(
            Department.id == current_user.department_id,
        )

    return db.session.scalars(query).all()


def get_publication_types():
    return db.session.scalars(
        select(ReferenceValue)
        .where(
            ReferenceValue.type
            == PUBLICATION_TYPE_REFERENCE,
            ReferenceValue.is_active.is_(True),
        )
        .order_by(
            ReferenceValue.sort_order,
            ReferenceValue.name_ru,
        )
    ).all()


def get_plan_item_query(
    year=None,
    department_id=None,
    publication_type_id=None,
):
    query = (
        select(
            PlanItem.id.label(
                "plan_item_id"
            ),
            Plan.year.label(
                "year"
            ),
            Department.id.label(
                "department_id"
            ),
            Department.name.label(
                "department_name"
            ),
            ReferenceValue.id.label(
                "publication_type_id"
            ),
            ReferenceValue.name_ru.label(
                "publication_type_name_ru"
            ),
            ReferenceValue.name_kk.label(
                "publication_type_name_kk"
            ),
            Publication.id.label(
                "publication_id"
            ),
            Publication.verification_status.label(
                "verification_status"
            ),
        )
        .select_from(PlanItem)
        .join(
            Plan,
            Plan.id == PlanItem.plan_id,
        )
        .join(
            Department,
            Department.id == Plan.department_id,
        )
        .join(
            ReferenceValue,
            ReferenceValue.id
            == PlanItem.publication_type_id,
        )
        .outerjoin(
            Publication,
            Publication.plan_item_id
            == PlanItem.id,
        )
        .where(
            Plan.status == ANALYTICS_PLAN_STATUS,
        )
    )

    if year is not None:
        query = query.where(
            Plan.year == year,
        )

    if department_id is not None:
        query = query.where(
            Department.id == department_id,
        )

    if publication_type_id is not None:
        query = query.where(
            PlanItem.publication_type_id
            == publication_type_id,
        )

    return query


def get_summary_by_publication_type(filters):
    base_query = get_plan_item_query(
        **filters,
    ).subquery()

    actual_expression = case(
        (
            base_query.c.verification_status
            == ANALYTICS_FACT_STATUS,
            1,
        ),
        else_=0,
    )

    rows = db.session.execute(
        select(
            base_query.c.publication_type_id,
            base_query.c.publication_type_name_ru,
            base_query.c.publication_type_name_kk,
            func.count(
                base_query.c.plan_item_id
            ).label(
                "planned_count"
            ),
            func.sum(
                actual_expression
            ).label(
                "actual_count"
            ),
        )
        .group_by(
            base_query.c.publication_type_id,
            base_query.c.publication_type_name_ru,
            base_query.c.publication_type_name_kk,
        )
        .order_by(
            base_query.c.publication_type_name_ru,
        )
    ).all()

    result = []

    for row in rows:
        planned_count = int(
            row.planned_count or 0
        )
        actual_count = int(
            row.actual_count or 0
        )

        result.append(
            {
                "publication_type_id": (
                    row.publication_type_id
                ),
                "name_ru": (
                    row.publication_type_name_ru
                ),
                "name_kk": (
                    row.publication_type_name_kk
                ),
                "planned_count": planned_count,
                "actual_count": actual_count,
                "progress": calculate_percentage(
                    planned_count,
                    actual_count,
                ),
            }
        )

    return result


def get_summary_by_department(filters):
    base_query = get_plan_item_query(
        **filters,
    ).subquery()

    actual_expression = case(
        (
            base_query.c.verification_status
            == ANALYTICS_FACT_STATUS,
            1,
        ),
        else_=0,
    )

    rows = db.session.execute(
        select(
            base_query.c.department_id,
            base_query.c.department_name,
            func.count(
                base_query.c.plan_item_id
            ).label(
                "planned_count"
            ),
            func.sum(
                actual_expression
            ).label(
                "actual_count"
            ),
        )
        .group_by(
            base_query.c.department_id,
            base_query.c.department_name,
        )
        .order_by(
            base_query.c.department_name,
        )
    ).all()

    result = []

    for row in rows:
        planned_count = int(
            row.planned_count or 0
        )
        actual_count = int(
            row.actual_count or 0
        )

        result.append(
            {
                "department_id": row.department_id,
                "name": row.department_name,
                "planned_count": planned_count,
                "actual_count": actual_count,
                "progress": calculate_percentage(
                    planned_count,
                    actual_count,
                ),
            }
        )

    return result


def get_total_summary(filters):
    base_query = get_plan_item_query(
        **filters,
    ).subquery()

    actual_expression = case(
        (
            base_query.c.verification_status
            == ANALYTICS_FACT_STATUS,
            1,
        ),
        else_=0,
    )

    row = db.session.execute(
        select(
            func.count(
                base_query.c.plan_item_id
            ).label(
                "planned_count"
            ),
            func.sum(
                actual_expression
            ).label(
                "actual_count"
            ),
        )
    ).one()

    planned_count = int(
        row.planned_count or 0
    )
    actual_count = int(
        row.actual_count or 0
    )

    return {
        "planned_count": planned_count,
        "actual_count": actual_count,
        "progress": calculate_percentage(
            planned_count,
            actual_count,
        ),
    }


def build_analytics_data(filters):
    return {
        "total": get_total_summary(filters),
        "by_publication_type": (
            get_summary_by_publication_type(
                filters
            )
        ),
        "by_department": (
            get_summary_by_department(
                filters
            )
        ),
    }


@analytics_bp.route("/")
@roles_required(
    ROLE_ADMIN,
    ROLE_DEPARTMENT_HEAD,
)
def dashboard():
    filters = get_filters()

    data = build_analytics_data(
        filters
    )

    return render_template(
        "analytics/dashboard.html",
        summary=data["total"],
        by_publication_type=(
            data["by_publication_type"]
        ),
        by_department=data["by_department"],
        years=get_years(),
        departments=get_departments(),
        publication_types=get_publication_types(),
        selected_year=filters["year"],
        selected_department_id=(
            filters["department_id"]
        ),
        selected_publication_type_id=(
            filters["publication_type_id"]
        ),
    )


@analytics_bp.route("/export.xlsx")
@roles_required(
    ROLE_ADMIN,
    ROLE_DEPARTMENT_HEAD,
)
def export_excel():
    filters = get_filters()

    data = build_analytics_data(
        filters
    )

    workbook = Workbook()

    summary_sheet = workbook.active
    summary_sheet.title = translate("Итоги")

    summary_sheet.append(
        [
            translate("Показатель"),
            translate("Значение"),
        ]
    )

    summary_sheet.append(
        [
            translate("План"),
            data["total"]["planned_count"],
        ]
    )

    summary_sheet.append(
        [
            translate("Факт"),
            data["total"]["actual_count"],
        ]
    )

    summary_sheet.append(
        [
            translate("Выполнение, %"),
            data["total"]["progress"],
        ]
    )

    type_sheet = workbook.create_sheet(
        translate("По типам публикаций")
    )

    type_sheet.append(
        [
            translate("Тип публикации"),
            translate("План"),
            translate("Факт"),
            translate("Выполнение, %"),
        ]
    )

    for item in data["by_publication_type"]:
        type_sheet.append(
            [
                (item["name_kk"] if current_language() == "kk" else item["name_ru"])
                or item["name_ru"]
                or item["name_kk"]
                or "",
                item["planned_count"],
                item["actual_count"],
                item["progress"],
            ]
        )

    department_sheet = workbook.create_sheet(
        translate("По кафедрам")
    )

    department_sheet.append(
        [
            translate("Кафедра"),
            translate("План"),
            translate("Факт"),
            translate("Выполнение, %"),
        ]
    )

    for item in data["by_department"]:
        department_sheet.append(
            [
                item["name"],
                item["planned_count"],
                item["actual_count"],
                item["progress"],
            ]
        )

    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = (
            worksheet.dimensions
        )

        for column_cells in worksheet.columns:
            maximum = max(
                (
                    len(str(cell.value))
                    for cell in column_cells
                    if cell.value is not None
                ),
                default=0,
            )

            letter = get_column_letter(
                column_cells[0].column
            )

            worksheet.column_dimensions[
                letter
            ].width = min(
                max(maximum + 2, 12),
                60,
            )

    output = BytesIO()

    workbook.save(output)
    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name=(
            "publication_analytics.xlsx"
        ),
        mimetype=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
    )


@analytics_bp.route("/export.pdf")
@roles_required(
    ROLE_ADMIN,
    ROLE_DEPARTMENT_HEAD,
)
def export_pdf():
    filters = get_filters()

    data = build_analytics_data(
        filters
    )

    summary = data["total"]
    by_publication_type = (
        data["by_publication_type"]
    )
    by_department = data["by_department"]

    output = BytesIO()

    font_candidates = [
        Path("C:/Windows/Fonts/arial.ttf"),
        Path("C:/Windows/Fonts/calibri.ttf"),
        Path(
            "/usr/share/fonts/truetype/dejavu/"
            "DejaVuSans.ttf"
        ),
    ]

    font_path = next(
        (
            candidate
            for candidate in font_candidates
            if candidate.exists()
        ),
        None,
    )

    if font_path is not None:
        font_name = "AnalyticsUnicode"

        if (
            font_name
            not in pdfmetrics.getRegisteredFontNames()
        ):
            pdfmetrics.registerFont(
                TTFont(
                    font_name,
                    str(font_path),
                )
            )
    else:
        font_name = "Helvetica"

    document = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        leftMargin=8 * mm,
        rightMargin=8 * mm,
        topMargin=8 * mm,
        bottomMargin=8 * mm,
    )

    text_style = ParagraphStyle(
        "AnalyticsText",
        fontName=font_name,
        fontSize=7,
        leading=8,
    )

    header_style = ParagraphStyle(
        "AnalyticsHeader",
        parent=text_style,
        alignment=TA_CENTER,
    )

    title_style = ParagraphStyle(
        "AnalyticsTitle",
        fontName=font_name,
        fontSize=15,
        leading=18,
        spaceAfter=5 * mm,
    )

    section_style = ParagraphStyle(
        "AnalyticsSection",
        fontName=font_name,
        fontSize=10,
        leading=12,
        spaceAfter=2 * mm,
    )

    story = [
        Paragraph(
            translate("Аналитика публикационной деятельности"),
            title_style,
        )
    ]

    summary_table = Table(
        [
            [
                Paragraph(
                    translate("Показатель"),
                    header_style,
                ),
                Paragraph(
                    translate("Значение"),
                    header_style,
                ),
            ],
            [
                Paragraph(
                    translate("План"),
                    text_style,
                ),
                Paragraph(
                    str(
                        summary[
                            "planned_count"
                        ]
                    ),
                    text_style,
                ),
            ],
            [
                Paragraph(
                    translate("Факт"),
                    text_style,
                ),
                Paragraph(
                    str(
                        summary[
                            "actual_count"
                        ]
                    ),
                    text_style,
                ),
            ],
            [
                Paragraph(
                    translate("Выполнение"),
                    text_style,
                ),
                Paragraph(
                    f'{summary["progress"]}%',
                    text_style,
                ),
            ],
        ],
        colWidths=[
            55 * mm,
            30 * mm,
        ],
    )

    summary_table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.lightgrey,
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.3,
                    colors.grey,
                ),
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "TOP",
                ),
            ]
        )
    )

    story.append(summary_table)
    story.append(
        Spacer(
            1,
            5 * mm,
        )
    )

    story.append(
        Paragraph(
            translate("По типам публикаций"),
            section_style,
        )
    )

    type_rows = [
        [
            Paragraph(
                translate("Тип публикации"),
                header_style,
            ),
            Paragraph(
                translate("План"),
                header_style,
            ),
            Paragraph(
                translate("Факт"),
                header_style,
            ),
            Paragraph(
                "%",
                header_style,
            ),
        ]
    ]

    for item in by_publication_type:
        type_rows.append(
            [
                Paragraph(
                    (item["name_kk"] if current_language() == "kk" else item["name_ru"])
                    or item["name_ru"]
                    or item["name_kk"]
                    or "",
                    text_style,
                ),
                Paragraph(
                    str(
                        item[
                            "planned_count"
                        ]
                    ),
                    text_style,
                ),
                Paragraph(
                    str(
                        item[
                            "actual_count"
                        ]
                    ),
                    text_style,
                ),
                Paragraph(
                    f'{item["progress"]}%',
                    text_style,
                ),
            ]
        )

    type_table = Table(
        type_rows,
        repeatRows=1,
        colWidths=[
            80 * mm,
            25 * mm,
            25 * mm,
            25 * mm,
        ],
    )

    type_table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.lightgrey,
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.3,
                    colors.grey,
                ),
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "TOP",
                ),
            ]
        )
    )

    story.append(type_table)
    story.append(
        Spacer(
            1,
            5 * mm,
        )
    )

    story.append(
        Paragraph(
            translate("По кафедрам"),
            section_style,
        )
    )

    department_rows = [
        [
            Paragraph(
                translate("Кафедра"),
                header_style,
            ),
            Paragraph(
                translate("План"),
                header_style,
            ),
            Paragraph(
                translate("Факт"),
                header_style,
            ),
            Paragraph(
                "%",
                header_style,
            ),
        ]
    ]

    for item in by_department:
        department_rows.append(
            [
                Paragraph(
                    item["name"],
                    text_style,
                ),
                Paragraph(
                    str(
                        item[
                            "planned_count"
                        ]
                    ),
                    text_style,
                ),
                Paragraph(
                    str(
                        item[
                            "actual_count"
                        ]
                    ),
                    text_style,
                ),
                Paragraph(
                    f'{item["progress"]}%',
                    text_style,
                ),
            ]
        )

    department_table = Table(
        department_rows,
        repeatRows=1,
        colWidths=[
            105 * mm,
            25 * mm,
            25 * mm,
            25 * mm,
        ],
    )

    department_table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.lightgrey,
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.3,
                    colors.grey,
                ),
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "TOP",
                ),
            ]
        )
    )

    story.append(department_table)

    document.build(story)

    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name=(
            "publication_analytics.pdf"
        ),
        mimetype="application/pdf",
    )
